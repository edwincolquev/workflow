"""Plantilla Excel de repuestos y carga atómica de sus cambios."""

from __future__ import annotations

from io import BytesIO
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter

from catalogo.database import DEFAULT_DB_PATH, get_connection
from catalogo.descriptions import refresh_long_description
from catalogo.gestion import ESTADOS, NUMERICOS, guardar_repuesto, numero_gestion, normalizar_estado, validar_imagen
from catalogo.hoja_trabajo import CAMPOS_EDITABLES, LABELS, _validar_listas, valor_comparable

HEADERS = ("ID", *(LABELS[field] for field in CAMPOS_EDITABLES), "Fecha cambio de estado")
LEGACY_FIELDS = tuple(field for field in CAMPOS_EDITABLES if field != "imagen")
LEGACY_HEADERS = ("ID", *(LABELS[field] for field in LEGACY_FIELDS), "Fecha cambio de estado")
BASE_SHEET = "Originales"
DATA_SHEET = "Repuestos"


def _plain(field, value):
    if value is None or value == "":
        return None
    if field == "estado_gestion":
        return normalizar_estado(value)
    if field == "activo":
        if isinstance(value, bool) or value in (0, 1):
            return bool(value)
        text = str(value).strip().lower()
        if text in ("sí", "si", "s", "true", "1", "activo"):
            return True
        if text in ("no", "n", "false", "0", "inactivo"):
            return False
        raise ValueError("Activo debe ser Sí o No")
    if field in NUMERICOS:
        return numero_gestion(field, value)
    if field == "imagen":
        return validar_imagen(value)
    if not isinstance(value, (str, int, float)):
        raise ValueError(f"{LABELS[field]} debe ser texto")
    return str(value).strip() or None


def _export_value(field, value):
    if field == "estado_gestion":
        return ESTADOS[value]
    if field == "activo":
        return "Sí" if value else "No"
    return value


def plantilla_hoja_bytes(db_path: str | Path = DEFAULT_DB_PATH, *, marca_id=None, estado=None) -> bytes:
    """Exporta los repuestos de la marca y el estado seleccionados."""
    conditions, params = [], []
    if marca_id == -1:
        conditions.append("marca_id IS NULL")
    elif marca_id is not None:
        conditions.append("marca_id = ?")
        params.append(marca_id)
    if estado is not None:
        conditions.append("estado_gestion = ?")
        params.append(normalizar_estado(estado))
    where = " WHERE " + " AND ".join(conditions) if conditions else ""
    with get_connection(db_path) as connection:
        rows = connection.execute("SELECT * FROM vista_repuestos" + where + " ORDER BY id", params).fetchall()
    book = Workbook()
    sheet = book.active
    sheet.title = DATA_SHEET
    sheet.append(HEADERS)
    baseline = book.create_sheet(BASE_SHEET)
    baseline.append(HEADERS)
    for part in rows:
        state_date = part["fecha_actualizacion_estado"]
        values = [part["id"], *(_export_value(field, part[field]) for field in CAMPOS_EDITABLES),
                  str(state_date) if state_date is not None else None]
        sheet.append(values)
        baseline.append(values)
        for target in (sheet, baseline):
            for cell in target[target.max_row]:
                if isinstance(cell.value, str) and cell.value.startswith("="):
                    cell.data_type = "s"
    baseline.sheet_state = "hidden"
    sheet.freeze_panes = "C2"
    sheet.auto_filter.ref = f"A1:{get_column_letter(len(HEADERS))}{max(2, sheet.max_row)}"
    sheet.sheet_view.showGridLines = False
    sheet.row_dimensions[1].height = 30
    widths = [10, 22, 18, 45, 22, 24, 18, 22, 25, 18, 18, 18, 20, 12, 45, 27]
    for index, width in enumerate(widths, 1):
        sheet.column_dimensions[sheet.cell(1, index).column_letter].width = width
    for cell in sheet[1]:
        cell.fill = PatternFill("solid", fgColor="17365D")
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    sheet["A1"].comment = Comment("No cambie el ID de las filas existentes. Déjelo vacío para agregar un repuesto.", "Catálogo")
    sheet["B1"].comment = Comment("Código obligatorio y único. Las columnas de códigos están formateadas como texto.", "Catálogo")
    sheet["I1"].comment = Comment("Puede elegir cualquier estado, sin seguir una secuencia.", "Catálogo")
    sheet["O1"].comment = Comment("URL HTTP o HTTPS de la foto referencial. Puede dejarla vacía.", "Catálogo")
    sheet["P1"].comment = Comment("Fecha del último cambio de estado. La calcula la aplicación; no la edite.", "Catálogo")
    for row in sheet.iter_rows(min_row=2):
        row[0].font = Font(color="666666")
        row[-1].font = Font(color="666666")
        for cell in row[1:]:
            cell.fill = PatternFill("solid", fgColor="F1F7FD" if cell.row % 2 == 0 else "FFFFFF")
        for column in (2, 3):
            row[column - 1].number_format = "@"
        for column in (10, 11, 12, 13):
            row[column - 1].number_format = "0.0000"
    for row_number in range(sheet.max_row + 1, sheet.max_row + 201):
        sheet.cell(row_number, 2).number_format = "@"
        sheet.cell(row_number, 3).number_format = "@"
    state_list = DataValidation(type="list", formula1='"' + ",".join(ESTADOS.values()) + '"')
    state_list.error = "Seleccione un estado de la lista"
    state_list.showErrorMessage = True
    active_list = DataValidation(type="list", formula1='"Sí,No"')
    active_list.showErrorMessage = True
    sheet.add_data_validation(state_list)
    sheet.add_data_validation(active_list)
    state_list.add("I2:I1048576")
    active_list.add("N2:N1048576")
    output = BytesIO()
    book.save(output)
    return output.getvalue()


def _read_template(content: bytes):
    try:
        book = load_workbook(BytesIO(content), read_only=False, data_only=False)
    except Exception as error:
        raise ValueError(f"No se pudo abrir el Excel: {error}") from error
    if DATA_SHEET not in book or BASE_SHEET not in book:
        raise ValueError("Use la plantilla descargada desde Hoja de trabajo")
    sheet, baseline = book[DATA_SHEET], book[BASE_SHEET]
    headers = tuple(cell.value for cell in sheet[1][:len(HEADERS)])
    variants = {
        HEADERS: (CAMPOS_EDITABLES, True),
        HEADERS[:-1]: (CAMPOS_EDITABLES, False),
        LEGACY_HEADERS: (LEGACY_FIELDS, True),
        LEGACY_HEADERS[:-1]: (LEGACY_FIELDS, False),
    }
    if headers not in variants:
        raise ValueError("Los encabezados de Repuestos fueron modificados")
    if tuple(cell.value for cell in baseline[1][:len(HEADERS)]) != headers:
        raise ValueError("La referencia interna de la plantilla fue modificada")
    template_fields, has_state_date = variants[headers]

    originals = {}
    original_dates = {}
    for row in baseline.iter_rows(min_row=2, max_col=len(headers), values_only=True):
        if row[0] is None:
            continue
        part_id = _id(row[0])
        if part_id in originals:
            raise ValueError(f"ID {part_id} repetido en la referencia interna")
        originals[part_id] = {field: _plain(field, row[index + 1])
                              for index, field in enumerate(template_fields)}
        if has_state_date:
            original_dates[part_id] = row[-1]

    seen_ids, seen_codes, updates, inserts = set(), set(), [], []
    for number, cells in enumerate(sheet.iter_rows(min_row=2, max_col=len(headers)), 2):
        row = tuple(cell.value for cell in cells)
        if not any(value is not None and value != "" for value in row):
            continue
        if any(cell.data_type == "f" for cell in cells):
            raise ValueError(f"Fila {number}: las fórmulas no se admiten")
        try:
            part_id = _id(row[0]) if row[0] is not None else None
            data = {field: _plain(field, row[index + 1]) for index, field in enumerate(template_fields)}
            code = data["codigo_proveedor"]
            if not code:
                raise ValueError("Código proveedor es obligatorio")
            code_key = code.upper()
            if code_key in seen_codes:
                raise ValueError(f"Código proveedor {code} repetido en el archivo")
            seen_codes.add(code_key)
            if part_id is None:
                if has_state_date and row[-1] is not None:
                    raise ValueError("Fecha cambio de estado debe quedar vacía en repuestos nuevos")
                inserts.append((number, data))
            else:
                if part_id not in originals:
                    raise ValueError(f"ID {part_id} no pertenece a esta plantilla")
                if part_id in seen_ids:
                    raise ValueError(f"ID {part_id} repetido")
                if has_state_date and row[-1] != original_dates[part_id]:
                    raise ValueError("Fecha cambio de estado es automática y no se puede editar")
                seen_ids.add(part_id)
                fields = {field: value for field, value in data.items()
                          if valor_comparable(field, value) != valor_comparable(field, originals[part_id][field])}
                if fields:
                    updates.append((number, part_id, fields, code))
        except ValueError as error:
            raise ValueError(f"Fila {number}: {error}") from error
    return originals, original_dates, updates, inserts, template_fields


def _id(value):
    if isinstance(value, bool) or not isinstance(value, (int, float)) or value != int(value) or value <= 0:
        raise ValueError("ID inválido; deje la celda vacía para un repuesto nuevo")
    return int(value)


def revisar_plantilla_hoja(content: bytes):
    """Devuelve solo las filas que se guardarán; no toca la base."""
    _, _, updates, inserts, _ = _read_template(content)
    return updates, inserts


def _guardar_fila_excel(connection, part_id, current, fields):
    """Aísla la edición comercial En SAP en la función privada de Postgres."""
    if not (hasattr(connection, "raw") and current["estado_gestion"] == "en_sap"
            and set(fields) & set(NUMERICOS)):
        return guardar_repuesto(connection, fields, part_id, permitir_comercial_excel=True)[1]
    commercial = {field: fields[field] for field in NUMERICOS if field in fields}
    other = {field: value for field, value in fields.items() if field not in NUMERICOS}
    if other:
        guardar_repuesto(connection, other, part_id)
    latest = connection.execute(
        "SELECT minimo_compra, precio_compra, precio_venta, cantidad_compra FROM repuestos WHERE id = ?",
        (part_id,),
    ).fetchone()
    values = [commercial.get(field, latest[field]) for field in NUMERICOS]
    connection.execute(
        "SELECT catalogo_privado.actualizar_comercial_excel(?, ?, ?, ?, ?)",
        (part_id, *values),
    )
    return True


def importar_plantilla_hoja(content: bytes, db_path: str | Path = DEFAULT_DB_PATH):
    """Revalida el archivo y aplica todas las altas y ediciones en una transacción."""
    originals, original_dates, updates, inserts, template_fields = _read_template(content)
    with get_connection(db_path) as connection:
        connection.execute("BEGIN IMMEDIATE")
        for number, part_id, fields, _code in updates:
            try:
                if hasattr(connection, "raw"):
                    connection.execute("SELECT id FROM repuestos WHERE id = ? FOR UPDATE", (part_id,))
                current = connection.execute("SELECT * FROM vista_repuestos WHERE id = ?", (part_id,)).fetchone()
                if current is None:
                    raise ValueError("el repuesto fue eliminado")
                if any(valor_comparable(field, current[field]) != valor_comparable(field, originals[part_id][field])
                       for field in template_fields):
                    raise ValueError("el repuesto cambió desde la descarga; descargue una plantilla nueva")
                if part_id in original_dates and str(current["fecha_actualizacion_estado"]) != original_dates[part_id]:
                    raise ValueError("el estado cambió desde la descarga; descargue una plantilla nueva")
                _validar_listas(connection, current, fields)
                changed = _guardar_fila_excel(connection, part_id, current, fields)
                if changed and set(fields) & {"descripcion", "medida"}:
                    refresh_long_description(connection, part_id)
            except Exception as error:
                raise ValueError(f"Fila {number}, ID {part_id}: {error}") from error
        for number, data in inserts:
            try:
                code = data["codigo_proveedor"]
                exists = connection.execute(
                    "SELECT id FROM repuestos WHERE upper(trim(codigo_proveedor)) = upper(trim(?))", (code,)
                ).fetchone()
                if exists:
                    raise ValueError(f"Código proveedor {code} ya existe (ID {exists[0]})")
                if data["activo"] is None:
                    data["activo"] = True
                created_id, _ = guardar_repuesto(connection, data, permitir_comercial_excel=True)
                refresh_long_description(connection, created_id)
            except Exception as error:
                raise ValueError(f"Fila {number}: {error}") from error
    return len(updates), len(inserts)
