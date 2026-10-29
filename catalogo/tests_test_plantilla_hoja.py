from io import BytesIO
from contextlib import contextmanager
import re

from openpyxl import load_workbook

from catalogo.database import get_connection, initialize_database
from catalogo.gestion import guardar_repuesto
from catalogo.plantilla_hoja import (
    HEADERS, importar_plantilla_hoja, plantilla_hoja_bytes, revisar_plantilla_hoja,
)


@contextmanager
def raises(message):
    try:
        yield
    except ValueError as error:
        assert re.search(message, str(error)), str(error)
    else:
        raise AssertionError("La operación debió rechazar la plantilla")


def _catalogo(tmp_path):
    path = tmp_path / "plantilla.sqlite3"
    initialize_database(path)
    with get_connection(path) as connection:
        guardar_repuesto(connection, {"codigo_proveedor": "001", "descripcion": "Original", "precio_compra": 2})
        guardar_repuesto(connection, {"codigo_proveedor": "002", "descripcion": "Sin cambios"})
    return path


def _edit(content, edit):
    book = load_workbook(BytesIO(content))
    edit(book["Repuestos"])
    output = BytesIO()
    book.save(output)
    return output.getvalue()


def test_excel_updates_and_adds_without_touching_unchanged_rows(tmp_path):
    path = _catalogo(tmp_path)
    content = plantilla_hoja_bytes(path)
    book = load_workbook(BytesIO(content))
    assert tuple(cell.value for cell in book["Repuestos"][1]) == HEADERS
    assert book["Originales"].sheet_state == "hidden"

    def edit(sheet):
        sheet["D2"] = "Actualizada"
        sheet["K2"] = 0
        sheet.append([None, "003", None, "Nuevo", None, None, None, None,
                      "Propuesto", None, None, None, None, "Sí"])

    content = _edit(content, edit)
    updates, inserts = revisar_plantilla_hoja(content)
    assert len(updates) == len(inserts) == 1
    assert importar_plantilla_hoja(content, path) == (1, 1)
    with get_connection(path) as connection:
        rows = connection.execute("SELECT codigo_proveedor, descripcion, precio_compra, descripcion_larga "
                                  "FROM repuestos ORDER BY codigo_proveedor").fetchall()
    assert [(row[0], row[1], row[2]) for row in rows] == [
        ("001", "Actualizada", 0), ("002", "Sin cambios", None), ("003", "Nuevo", None)
    ]
    assert rows[0][3] == "Actualizada"


def test_excel_conflict_cancels_entire_import(tmp_path):
    path = _catalogo(tmp_path)
    content = _edit(plantilla_hoja_bytes(path), lambda sheet: (
        setattr(sheet["D2"], "value", "Desde Excel"),
        sheet.append([None, "003", None, "Nuevo"]),
    ))
    with get_connection(path) as connection:
        guardar_repuesto(connection, {"descripcion": "Desde otra sesión"}, 1)
    with raises("cambió desde la descarga"):
        importar_plantilla_hoja(content, path)
    with get_connection(path) as connection:
        assert connection.execute("SELECT COUNT(*) FROM repuestos").fetchone()[0] == 2
        assert connection.execute("SELECT descripcion FROM repuestos WHERE id = 1").fetchone()[0] == "Desde otra sesión"


def test_excel_invalid_row_rolls_back_previous_updates(tmp_path):
    path = _catalogo(tmp_path)
    def edit(sheet):
        sheet["D2"] = "No guardar"
        sheet["I3"] = "Pedido"
    content = _edit(plantilla_hoja_bytes(path), edit)
    with raises("Fila 3"):
        importar_plantilla_hoja(content, path)
    with get_connection(path) as connection:
        assert connection.execute("SELECT descripcion FROM repuestos WHERE id = 1").fetchone()[0] == "Original"


def test_excel_rejects_changed_header_and_duplicate_code(tmp_path):
    path = _catalogo(tmp_path)
    original = plantilla_hoja_bytes(path)
    with raises("encabezados"):
        revisar_plantilla_hoja(_edit(original, lambda sheet: setattr(sheet["B1"], "value", "Proveedor")))
    with raises("repetido"):
        revisar_plantilla_hoja(_edit(original, lambda sheet: sheet.append([None, "001", None, "Duplicado"])))
